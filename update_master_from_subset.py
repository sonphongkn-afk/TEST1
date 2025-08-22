#!/usr/bin/env python3

import argparse
import os
import sys
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook


def normalize_col_map(columns: Iterable[object]) -> Dict[str, object]:
	return {str(c).strip().lower(): c for c in columns}


def find_stt_column(df: pd.DataFrame, preferred: Optional[str] = None) -> str:
	if preferred is not None:
		# Try exact name, else case-insensitive
		if preferred in df.columns:
			return preferred
		norm = normalize_col_map(df.columns)
		key = str(preferred).strip().lower()
		if key in norm:
			return norm[key]
		raise ValueError(f"Không tìm thấy cột STT '{preferred}' trong file.")

	norm = normalize_col_map(df.columns)
	candidates = [
		"số thứ tự", "so thu tu", "stt", "index", "id", "mã", "ma",
	]
	for k in candidates:
		if k in norm:
			return norm[k]
	raise ValueError("Không tìm thấy cột STT. Hãy chỉ định bằng --stt_col_name.")


def detect_question_columns(df: pd.DataFrame, num_questions: int, first_index: Optional[int]) -> List[str]:
	cols = list(df.columns)
	if first_index is not None:
		end = first_index + num_questions
		if end > len(cols):
			raise ValueError(
				f"first_question_col_index + num_questions vượt quá số cột. {first_index}+{num_questions}>{len(cols)}"
			)
		return cols[first_index:end]
	if len(cols) < num_questions:
		raise ValueError(
			f"Bảng có {len(cols)} cột, ít hơn số câu hỏi ({num_questions}). Hãy kiểm tra tham số."
		)
	return cols[-num_questions:]


def normalize_key(value: object) -> Optional[str]:
	if pd.isna(value):
		return None
	s = str(value).strip()
	if s == "":
		return None
	# Nếu là số (vd 1.0, 1), chuyển về chuỗi số nguyên để tăng khả năng khớp
	try:
		f = float(s.replace(",", ".")) if any(ch in s for ch in ",.") else float(s)
		if np.isfinite(f):
			iv = int(f)
			if abs(f - iv) < 1e-9:
				return str(iv)
	except Exception:
		pass
	return s


def choose_subset_question_columns(
	sub_df: pd.DataFrame,
	master_question_cols: Sequence[str],
	sub_first_index: Optional[int],
	questions: int,
) -> List[str]:
	# Ưu tiên map theo tên cột trùng khớp
	matched_by_name = [c for c in master_question_cols if c in sub_df.columns]
	if len(matched_by_name) == questions:
		return matched_by_name

	# Nếu không đủ theo tên, dùng chỉ số nếu được cung cấp
	if sub_first_index is not None:
		cols = list(sub_df.columns)
		end = sub_first_index + questions
		if end > len(cols):
			raise ValueError(
				f"subset_first_question_col_index + questions vượt quá số cột subset. {sub_first_index}+{questions}>{len(cols)}"
			)
		return cols[sub_first_index:end]

	# Fallback: dùng 51 cột cuối cùng của subset
	cols = list(sub_df.columns)
	if len(cols) < questions:
		raise ValueError(
			f"Subset có {len(cols)} cột, ít hơn số câu hỏi ({questions}). Hãy truyền --subset_first_question_col_index."
		)
	return cols[-questions:]


def update_master_from_subset(
	master_path: str,
	subset_path: str,
	output_path: Optional[str],
	questions: int,
	master_first_idx: Optional[int],
	subset_first_idx: Optional[int],
	stt_col_name: Optional[str],
	master_sheet: Optional[str],
	subset_sheet: Optional[str],
) -> Tuple[str, int, int, int, int, int, int]:
	# Đọc vào DataFrame để xác định cột và map dữ liệu subset
	master_df = pd.read_excel(master_path, engine="openpyxl", sheet_name=master_sheet if master_sheet else 0)
	sub_df = pd.read_excel(subset_path, engine="openpyxl", sheet_name=subset_sheet if subset_sheet else 0)

	# Xác định cột STT cho cả master và subset (theo tên, không phân biệt hoa/thường)
	stt_master = find_stt_column(master_df, preferred=stt_col_name)
	stt_subset = find_stt_column(sub_df, preferred=stt_col_name)

	# Xác định 51 cột câu hỏi ở master và subset
	master_q_cols = detect_question_columns(master_df, questions, master_first_idx)
	subset_q_cols = choose_subset_question_columns(sub_df, master_q_cols, subset_first_idx, questions)

	# Map từ STT -> Series dữ liệu câu hỏi ở subset
	subset_map: Dict[str, pd.Series] = {}
	for _, row in sub_df.iterrows():
		key = normalize_key(row.get(stt_subset))
		if key is None:
			continue
		subset_map[key] = row[subset_q_cols]

	# Mở workbook gốc bằng openpyxl để GIỮ NGUYÊN mọi sheet/định dạng; chỉ cập nhật ô cần thiết
	wb = load_workbook(master_path)
	ws = wb[master_sheet] if master_sheet else wb[wb.sheetnames[0]]

	# Tạo map tiêu đề hàng 1 -> chỉ số cột (1-based)
	header_cells = list(ws.iter_rows(min_row=1, max_row=1, values_only=False))[0]
	header_map: Dict[str, int] = {}
	for col_idx, cell in enumerate(header_cells, start=1):
		val = cell.value
		key = str(val).strip().lower() if val is not None else ""
		if key and key not in header_map:
			header_map[key] = col_idx

	# Xác định vị trí cột STT trong sheet
	stt_col_idx = header_map.get(str(stt_master).strip().lower())
	if stt_col_idx is None:
		# Fallback theo vị trí cột trong DataFrame
		try:
			stt_pos = list(master_df.columns).index(stt_master)
			stt_col_idx = stt_pos + 1
		except Exception:
			raise ValueError("Không xác định được cột STT trong sheet master.")

	# Vị trí các cột câu hỏi trong sheet
	df_cols = list(master_df.columns)
	master_q_col_indices: List[int] = []
	for mcol in master_q_cols:
		key = str(mcol).strip().lower()
		cidx = header_map.get(key)
		if cidx is None:
			# Fallback theo vị trí trong DataFrame
			try:
				pos = df_cols.index(mcol)
				cidx = pos + 1
			except Exception:
				raise ValueError(f"Không xác định được vị trí cột '{mcol}' trong sheet master.")
		master_q_col_indices.append(cidx)

	# Cập nhật theo từng hàng khớp STT (giữ nguyên hàng/ô khác)
	matched_rows = 0
	updated_rows = 0
	updated_cells = 0
	equal_rows = 0
	nan_only_rows = 0
	not_matched = 0

	max_row = ws.max_row
	for r in range(2, max_row + 1):
		stt_val = ws.cell(row=r, column=stt_col_idx).value
		key = normalize_key(stt_val)
		if key is None or key not in subset_map:
			not_matched += 1
			continue
		matched_rows += 1
		new_vals_series = subset_map[key]
		row_updates = 0
		has_any_non_na = False
		for cidx, scol in zip(master_q_col_indices, subset_q_cols):
			new_val = new_vals_series.get(scol)
			if pd.isna(new_val):
				# Không ghi đè bằng giá trị trống
				continue
			has_any_non_na = True
			# Chuẩn hóa giá trị 1..5 về int nếu hợp lệ
			try:
				iv = int(float(new_val))
				if 1 <= iv <= 5:
					write_val = iv
				else:
					write_val = new_val
			except Exception:
				write_val = new_val

			old_val = ws.cell(row=r, column=cidx).value
			if old_val != write_val:
				ws.cell(row=r, column=cidx).value = write_val
				row_updates += 1
		if row_updates > 0:
			updated_rows += 1
			updated_cells += row_updates
		else:
			if has_any_non_na:
				equal_rows += 1
			else:
				nan_only_rows += 1

	# Lưu workbook (giữ nguyên tất cả sheet/định dạng)
	if output_path is None:
		base, ext = os.path.splitext(os.path.basename(master_path))
		output_path = os.path.join(os.path.dirname(master_path), f"{base}_updated.xlsx")
	wb.save(output_path)

	return output_path, updated_rows, updated_cells, not_matched, matched_rows, equal_rows, nan_only_rows


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
	p = argparse.ArgumentParser(
		description=(
			"Cập nhật file Excel tổng hợp (master) bằng dữ liệu subset theo cột STT. "
			"Chỉ cập nhật các cột câu hỏi (mặc định 51 cột) cho các STT có trong subset."
		)
	)
	p.add_argument("--master", required=True, help="Đường dẫn file Excel master (tổng hợp)")
	p.add_argument("--subset", required=True, help="Đường dẫn file Excel subset (ví dụ file *_ket_qua.xlsx)")
	p.add_argument("--output", default=None, help="Đường dẫn file kết quả (mặc định: <master>_updated.xlsx)")
	p.add_argument("--questions", type=int, default=51, help="Số lượng cột câu hỏi (mặc định: 51)")
	p.add_argument("--master_first_question_col_index", type=int, default=None, help="Chỉ số cột bắt đầu khối câu hỏi ở master (0-based). Mặc định: dùng 51 cột cuối cùng")
	p.add_argument("--subset_first_question_col_index", type=int, default=None, help="Chỉ số cột bắt đầu khối câu hỏi ở subset (0-based). Nếu không set, script sẽ cố map theo tên cột, nếu không đủ sẽ dùng 51 cột cuối cùng")
	p.add_argument("--stt_col_name", type=str, default=None, help="Tên cột STT nếu muốn chỉ định rõ (nếu không, script sẽ tự dò: STT, Số thứ tự, v.v.)")
	p.add_argument("--master_sheet", type=str, default=None, help="Tên sheet ở master (mặc định: sheet đầu tiên)")
	p.add_argument("--subset_sheet", type=str, default="DATA_FILLED", help="Tên sheet ở subset (mặc định: DATA_FILLED)")
	return p.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
	args = parse_args(argv)
	try:
		output_path, updated_rows, updated_cells, not_matched, matched_rows, equal_rows, nan_only_rows = update_master_from_subset(
			master_path=args.master,
			subset_path=args.subset,
			output_path=args.output,
			questions=args.questions,
			master_first_idx=args.master_first_question_col_index,
			subset_first_idx=args.subset_first_question_col_index,
			stt_col_name=args.stt_col_name,
			master_sheet=args.master_sheet,
			subset_sheet=args.subset_sheet,
		)
		print(
			f"Đã tạo file cập nhật: {output_path}\n"
			f"- Hàng có STT khớp: {matched_rows}\n"
			f"- Hàng cập nhật (có thay đổi): {updated_rows}\n"
			f"- Hàng khớp nhưng không đổi (đã giống): {equal_rows}\n"
			f"- Hàng khớp nhưng subset toàn NaN ở 51 cột: {nan_only_rows}\n"
			f"- Ô cập nhật: {updated_cells}\n"
			f"- Hàng master không khớp STT: {not_matched}"
		)
		return 0
	except Exception as e:
		print(f"Lỗi: {e}")
		return 1


if __name__ == "__main__":
	sys.exit(main())